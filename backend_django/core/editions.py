"""Présentation des états déjà consultés : aucun calcul métier ni changement de statut."""
import io
import re
import math
from pathlib import Path
from datetime import datetime
from decimal import Decimal, InvalidOperation
from xml.sax.saxutils import escape
from django.http import HttpResponse
from django.utils.http import content_disposition_header
from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from .auth import assert_acces_societe
from .models import Societe
from .rh_views import ident
from .views import _societe_param

TEAL='176B61'; NAVY='183342'; GRAY='61717C'


def texte(v,limite=20000):
    if not isinstance(v,(str,int,float)) or isinstance(v,bool): raise ValidationError('Texte de document invalide.')
    if isinstance(v,float) and not math.isfinite(v): raise ValidationError('Nombre non fini.')
    s=str(v)
    if len(s)>limite or any(ord(c)<32 and c not in '\n\r\t' for c in s): raise ValidationError('Texte trop long ou caractère non imprimable.')
    return s


def normaliser(p):
    if not isinstance(p,dict): raise ValidationError('Document invalide.')
    titre=texte(p.get('titre','Document'),250);sous=texte(p.get('sous_titre',''),2000)
    blocks=p.get('blocs')
    if not isinstance(blocks,list) or not 1<=len(blocks)<=500: raise ValidationError('Document vide ou trop volumineux.')
    out=[];nb=0
    for b in blocks:
        if not isinstance(b,dict): raise ValidationError('Bloc invalide.')
        kind=b.get('type')
        if kind=='table':
            cols=b.get('colonnes',[]);rows=b.get('lignes',[])
            if not isinstance(cols,list) or not 1<=len(cols)<=24 or not isinstance(rows,list): raise ValidationError('Tableau invalide : maximum 24 colonnes.')
            nb+=len(rows)*len(cols)
            if nb>100000: raise ValidationError('Export trop volumineux. Réduisez la période ou les filtres.')
            if any(not isinstance(r,list) or len(r)!=len(cols) for r in rows): raise ValidationError('Colonnes incohérentes.')
            for r in rows:
                for v in r: texte(v if v is not None else '')
            out.append({'type':'table','colonnes':[texte(c,250) for c in cols],
                'lignes':[[v if isinstance(v,(int,float)) and not isinstance(v,bool) else texte(v if v is not None else '') for v in r] for r in rows]})
        elif kind in ['texte','titre','total','signatures']:
            out.append({'type':kind,'texte':texte(b.get('texte',''))})
        else: raise ValidationError('Type de bloc non autorisé.')
    return {'titre':titre,'sous_titre':sous,'blocs':out,'paysage':any(b['type']=='table' and len(b['colonnes'])>7 for b in out)}


def identite(s):
    return {'nom':s.nom,'code':s.code,'ville':s.ville or '',
            'references':' · '.join(f'{label} : {v}' for label,v in [('RCCM',s.rccm),('ID. NAT.',s.id_nat),('NIF',s.nif)] if v)}


@api_view(['GET'])
def societe(request):
    sid=ident(_societe_param(request));assert_acces_societe(request.user,sid)
    return Response(identite(Societe.objects.get(id=sid)))


def pdf(spec,soc):
    from reportlab.pdfgen import canvas
    from reportlab.lib import colors
    def couleur(v): return colors.HexColor("#"+v)
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT,TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab import __file__ as reportlab_file
    fonts=Path(reportlab_file).parent/'fonts'
    if 'Edition' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Edition',str(fonts/'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('EditionBold',str(fonts/'VeraBd.ttf')))
    def clean(v): return escape(str(v).replace('\u202f',' ').replace('\u00a0',' ')).replace('\n','<br/>')
    size=landscape(A4) if spec['paysage'] else A4
    output=io.BytesIO();width=size[0]-80
    normal=ParagraphStyle('Normal',fontName='Edition',fontSize=9,leading=13,textColor=couleur(NAVY),spaceAfter=7,splitLongWords=True)
    bold=ParagraphStyle('Titre',parent=normal,fontName='EditionBold',fontSize=11,leading=15,spaceBefore=10,spaceAfter=8)
    title=ParagraphStyle('GrandTitre',parent=bold,fontSize=20,leading=25,spaceBefore=0,spaceAfter=10)
    cell=ParagraphStyle('Cellule',parent=normal,fontSize=8,leading=11,spaceAfter=0)
    right=ParagraphStyle('Nombre',parent=cell,alignment=TA_RIGHT)
    cell_bold=ParagraphStyle('CelluleTotal',parent=cell,fontName='EditionBold',textColor=colors.HexColor('#176B61'))
    right_bold=ParagraphStyle('NombreTotal',parent=cell_bold,alignment=TA_RIGHT)
    th=ParagraphStyle('Entete',parent=cell,fontName='EditionBold',textColor=colors.white)
    total=ParagraphStyle('Total',parent=bold,textColor=couleur(TEAL),alignment=TA_RIGHT)
    def p(v,style=normal): return Paragraph(clean(v),style)
    doc=SimpleDocTemplate(output,pagesize=size,rightMargin=40,leftMargin=40,topMargin=106,bottomMargin=54,
        title=spec['titre'],author=soc['nom'],pageCompression=1)
    def page(c,d):
        c.saveState();w,h=size
        c.setFillColor(couleur(TEAL));c.rect(40,h-38,34,4,fill=1,stroke=0)
        company=p(soc['nom'],ParagraphStyle('Société',parent=bold,fontSize=13,leading=16,spaceBefore=0))
        _,ch=company.wrap(w-180,42);company.drawOn(c,40,h-50-ch)
        c.setFont('Edition',8);c.setFillColor(couleur(GRAY));c.drawRightString(w-40,h-48,datetime.now().strftime('%d/%m/%Y · %H:%M'))
        if soc['ville']: c.drawRightString(w-40,h-61,soc['ville'][:35])
        c.setStrokeColor(couleur('DCE5E5'));c.line(40,h-92,w-40,h-92)
        c.line(40,40,w-40,40)
        foot=p(soc['references'] or soc['nom'],ParagraphStyle('Footer',parent=normal,fontSize=6.5,leading=9))
        _,fh=foot.wrap(w-150,28);foot.drawOn(c,40,32-fh)
        c.setFont('Edition',7);c.drawRightString(w-40,23,f'Page {d.page}');c.restoreState()
    story=[p(spec['titre'],title)]
    if spec['sous_titre']: story.append(p(spec['sous_titre']))
    for b in spec['blocs']:
        if b['type']=='table':
            cols=b['colonnes'];n=len(cols)
            # Weight descriptive columns; all content wraps and headers repeat.
            weights=[2.4 if re.search(r'désignation|description|rubrique|détail|agent|libellé|nom',c,re.I) else 1 for c in cols]
            widths=[width*x/sum(weights) for x in weights]
            rows=[[p(x,th) for x in cols]]
            def affiche(v):
                if isinstance(v,(int,float)):
                    precision=max(2,min(8,-Decimal(str(v)).as_tuple().exponent))
                    return format(v,f',.{precision}f').replace(',',' ').replace('.',',')
                return str(v)
            for r in b['lignes']:
                fort=bool(r and re.match(r'^(total|net à payer|résultat|solde)',str(r[0]),re.I))
                rows.append([p(affiche(v),(right_bold if fort else right) if re.fullmatch(r'[-+\d\s.,%]+(?: USD| CDF| \$)?',str(v)) else (cell_bold if fort else cell)) for v in r])
            t=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT',splitByRow=1,splitInRow=1)
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),couleur(TEAL)),
                ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,couleur('F3F7F7')]),
                ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
                ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),
                ('LINEBELOW',(0,0),(-1,0),.7,couleur(TEAL)),('LINEBELOW',(0,1),(-1,-1),.25,couleur('DCE5E5'))]))
            story.extend([Spacer(1,7),t,Spacer(1,12)])
        elif b['type']=='signatures':
            labels=[t.strip() for t in b['texte'].split('|') if t.strip()] or ['Signature']
            signature_style=ParagraphStyle('Signature',parent=cell,fontName='EditionBold',fontSize=8.5,leading=12)
            st=Table([[p(t,signature_style) for t in labels],[p('Nom, date et signature',cell) for t in labels]],colWidths=[width/len(labels)]*len(labels))
            st.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,0),10),
                ('BOTTOMPADDING',(0,0),(-1,0),35),('LINEABOVE',(0,0),(-1,0),.5,colors.HexColor('#A0B4B5')),
                ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),16)]))
            story.append(KeepTogether([Spacer(1,26),st]))
        else: story.append(p(b['texte'],bold if b['type']=='titre' else total if b['type']=='total' else normal))
    doc.build(story,onFirstPage=page,onLaterPages=page)
    return output.getvalue()


def xlsx(spec,soc):
    # Runtime de l'ERP : exports natifs, sans dépendance à un service externe.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table,TableStyleInfo
    wb=Workbook();wb.remove(wb.active)
    tables=[b for b in spec['blocs'] if b['type']=='table']
    if not tables: tables=[{'colonnes':['Rubrique','Détail'],'lignes':[[b['type'],b['texte']] for b in spec['blocs']]}]
    notes=[b['texte'] for b in spec['blocs'] if b['type']!='table']
    def textcell(c,v): c.value=str(v);c.data_type='s'
    for ix,b in enumerate(tables,1):
        ws=wb.create_sheet('État' if len(tables)==1 else f'Tableau {ix}');cols=b['colonnes'];n=len(cols);last=get_column_letter(max(2,n))
        for row,value in [(1,soc['nom']),(2,spec['titre']),(3,spec['sous_titre']),(4,soc['references']),(5,'Édité le '+datetime.now().strftime('%d/%m/%Y à %H:%M'))]:
            ws.merge_cells(f'A{row}:{last}{row}');textcell(ws.cell(row,1),value)
            ws.cell(row,1).font=Font(name='Calibri',size=18 if row==2 else 12 if row==1 else 10,bold=row in [1,2],color=TEAL if row==2 else NAVY)
            ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=34 if row==2 else 27
        for j,c in enumerate(cols,1):
            textcell(ws.cell(7,j),c);ws.cell(7,j).font=Font(name='Calibri',bold=True,color='FFFFFF',size=10)
            ws.cell(7,j).fill=PatternFill('solid',fgColor=TEAL);ws.cell(7,j).alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[7].height=32
        for i,r in enumerate(b['lignes'],8):
            maxlines=1
            for j,v in enumerate(r,1):
                c=ws.cell(i,j);numeric=isinstance(v,(int,float))
                # Only amounts, quantities and rates are converted; identifiers keep leading zeros.
                if not numeric and re.search(r'montant|quantit|qté|qte|prix|p\.u\.|total|solde|débit|crédit|cump|valeur|coût|cout|net|brut|base|cnss|irpp|inpp|onem|\bht\b|\bttc\b|taux|^20\d{2}$',cols[j-1],re.I):
                    clean=str(v).replace('\u202f','').replace('\u00a0','').replace(' ','').replace(',','.')
                    if re.fullmatch(r'-?\d+(\.\d+)?',clean):
                        try:
                            num=Decimal(clean)
                            if num.is_finite() and len(num.as_tuple().digits)<=15: v=float(num);numeric=True
                        except InvalidOperation: pass
                if numeric: c.value=v;c.number_format='#,##0.00######;[Red](#,##0.00######);"–"'
                else: textcell(c,v)
                fort=bool(r and re.match(r'^(total|net à payer|résultat|solde)',str(r[0]),re.I))
                c.font=Font(name='Calibri',size=11,color=TEAL if fort else NAVY,bold=fort)
                c.alignment=Alignment(horizontal='right' if numeric else 'left',vertical='top',wrap_text=True)
                c.fill=PatternFill('solid',fgColor='F3F7F7' if i%2==0 else 'FFFFFF')
                maxlines=max(maxlines,len(str(v))//38+str(v).count('\n')+1)
            ws.row_dimensions[i].height=min(200,max(27,16*maxlines))
        end=7+len(b['lignes'])
        ws.auto_filter.ref=f'A7:{get_column_letter(n)}{max(7,end)}';ws.freeze_panes='A8';ws.sheet_view.showGridLines=False
        for j,c in enumerate(cols,1):
            samples=[str(r[j-1]) for r in b['lignes'][:500]]
            length=max([len(c),*(min(len(v),48) for v in samples)],default=16)
            ws.column_dimensions[get_column_letter(j)].width=max(15,min(50,length+3))
        largeur=sum(ws.column_dimensions[get_column_letter(j)].width for j in range(1,max(2,n)+1))
        ws.row_dimensions[7].height=max(32,max(math.ceil(len(c)/max(8,ws.column_dimensions[get_column_letter(j)].width-3))*14+10 for j,c in enumerate(cols,1)))
        for row in range(1,6):
            lignes=max(1,math.ceil(len(str(ws.cell(row,1).value))/max(12,largeur*.75)))
            ws.row_dimensions[row].height=max(ws.row_dimensions[row].height,lignes*(24 if row==2 else 15))
        cursor=end+3
        for note in notes:
            ws.merge_cells(start_row=cursor,start_column=1,end_row=cursor,end_column=max(2,n));c=ws.cell(cursor,1);textcell(c,note)
            c.font=Font(name='Calibri',size=10,color=GRAY);c.alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[cursor].height=max(30,min(180,len(note)//100*15+30));cursor+=1
        ws.print_title_rows='1:7';ws.print_options.horizontalCentered=True
        ws.page_setup.orientation='landscape' if n>7 else 'portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_area=f'A1:{last}{max(end,cursor-1)}';ws.oddFooter.right.text='Page &P / &N';ws.oddFooter.left.text=soc['nom'].replace('&','&&')[:100]
    wb.properties.creator=soc['nom'];wb.properties.title=spec['titre'];wb.properties.subject=spec['sous_titre']
    output=io.BytesIO();wb.save(output);return output.getvalue()


@api_view(['POST'])
def telecharger(request):
    sid=ident(_societe_param(request));assert_acces_societe(request.user,sid)
    if len(request.body)>4000000: raise ValidationError('Document trop volumineux.')
    spec=normaliser(request.data);soc=identite(Societe.objects.get(id=sid));fmt=request.data.get('format')
    if fmt not in ['pdf','xlsx']: raise ValidationError('Format non pris en charge.')
    data=pdf(spec,soc) if fmt=='pdf' else xlsx(spec,soc)
    r=HttpResponse(data,content_type='application/pdf' if fmt=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    nom=re.sub(r'[^\w .-]','',spec['titre'])[:100].strip() or 'Document'
    r['Content-Disposition']=content_disposition_header(True,f'{nom}.{fmt}');r['Cache-Control']='private, no-store'
    r['X-Content-Type-Options']='nosniff';return r
