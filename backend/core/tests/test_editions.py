import io
from django.test import TestCase
from rest_framework.test import APIClient
from core.tests import test_rh as base
from core.editions import normaliser,pdf,xlsx,identite
from openpyxl import load_workbook


class EditionsTests(TestCase):
    def setUp(self):
        base.RHTests.setUp(self)
        self.s.rccm='CD/LSH/RCCM-TEST';self.s.nif='NIF-TEST';self.s.save()
        self.spec={'titre':'Facture FA-2026-001','sous_titre':'Document fictif · USD · 15/09/2026',
            'blocs':[{'type':'texte','texte':'Client : Exemple & associés'},
                {'type':'table','colonnes':['Référence','Désignation','Quantité','Montant HT'],
                 'lignes':[['00123','Prestation de vérification',2,'1\u202f234,50'],['=HYPERLINK("https://evil.invalid")','=1+2',1,30]]},
                {'type':'total','texte':'Total TTC : 1 264,50 USD'},{'type':'signatures','texte':'Établi par | Réceptionné par'}]}

    def test_xlsx_types_accent_formules_et_impression(self):
        data=xlsx(normaliser(self.spec),identite(self.s));w=load_workbook(io.BytesIO(data));s=w.active
        self.assertEqual(s['A8'].value,'00123');self.assertEqual(s['A8'].data_type,'s')
        self.assertEqual(s['D8'].value,1234.5);self.assertEqual(s['A9'].data_type,'s');self.assertEqual(s['B9'].data_type,'s')
        self.assertEqual(s.freeze_panes,'A8');self.assertEqual(s.auto_filter.ref,'A7:D9')
        self.assertEqual(s.print_title_rows,'$1:$7');self.assertFalse(s.sheet_view.showGridLines)
        self.assertEqual(s['A1'].value,self.s.nom);self.assertIn('RCCM',s['A4'].value)

    def test_pdf_multipage_conserve_derniere_ligne_et_references(self):
        try: import fitz
        except ImportError: self.skipTest("PyMuPDF requis uniquement pour le contrôle visuel des PDF.")
        spec=normaliser({**self.spec,'blocs':[{'type':'table','colonnes':['Désignation','Montant USD'],
            'lignes':[[f'Ligne {i} — contrôle de réception et conformité des marchandises',i] for i in range(160)]}]})
        data=pdf(spec,identite(self.s));doc=fitz.open(stream=data,filetype='pdf')
        self.assertGreater(len(doc),2)
        texte=''.join(p.get_text() for p in doc);self.assertIn('Ligne 159',texte)
        for p in doc:
            self.assertIn('RCCM',p.get_text());self.assertIn('Page ',p.get_text());self.assertIn('Désignation',p.get_text())

    def test_exports_authentifies_et_societe_imposee(self):
        path=f'/api/editions/telecharger?societe_id={self.s.id}'
        self.assertEqual(APIClient().post(path,{**self.spec,'format':'pdf'},format='json').status_code,401)
        r=self.clients['DFI'].post(path,{**self.spec,'format':'xlsx','societe':{'nom':'Fausse société'}},format='json')
        self.assertEqual(r.status_code,200,r.content[:400]);self.assertIn('no-store',r['Cache-Control'])
        self.assertEqual(load_workbook(io.BytesIO(r.content)).active['A1'].value,self.s.nom)
        r=self.clients['DFI'].post(f'/api/editions/telecharger?societe_id={self.autre.id}',{**self.spec,'format':'pdf'},format='json')
        self.assertEqual(r.status_code,403)
        r=self.clients['DFI'].post(path,{**self.spec,'blocs':[{'type':'table','colonnes':['A'],'lignes':[['1','2']]}],'format':'pdf'},format='json')
        self.assertEqual(r.status_code,400)
